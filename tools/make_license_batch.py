#!/usr/bin/env python3
"""Creates a batch of Studiolo licence codes.

    python3 tools/make_license_batch.py --count 1000 --out "/path/outside/the/project"

* Writes the real codes ONLY to the --out folder (a spreadsheet and a CSV). Keep those files
  private: anyone who has them can activate Studiolo. Never put them in this repository.
* Adds only one-way fingerprints of the codes to brands/studiolo/license-codes.js, which ships
  with the app. The codes cannot be recovered from the fingerprints.

Each code is 16 random characters (80 bits) in Crockford base32, shown as XXXX-XXXX-XXXX-XXXX.
Fingerprint = first 16 bytes of SHA-512(batch salt + code), matching js/license.js.
"""
import argparse
import csv
import datetime
import hashlib
import json
import os
import re
import secrets

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CODES_JS = os.path.join(ROOT, 'brands', 'studiolo', 'license-codes.js')
ALPHABET = '0123456789ABCDEFGHJKMNPQRSTVWXYZ'


def new_code():
    raw = ''.join(secrets.choice(ALPHABET) for _ in range(16))
    return '-'.join(raw[i:i + 4] for i in range(0, 16, 4))


def fingerprint(salt, code):
    return hashlib.sha512((salt + code.replace('-', '')).encode()).hexdigest()[:32]


def load_batches():
    if not os.path.exists(CODES_JS):
        return []
    text = open(CODES_JS, encoding='utf-8').read()
    return json.loads(re.search(r'batches: (\[.*\])\s*\};', text, re.S).group(1))


def save_batches(batches):
    body = json.dumps(batches, indent=1)
    with open(CODES_JS, 'w', encoding='utf-8') as f:
        f.write('/*\n * Studiolo licence fingerprints, made by tools/make_license_batch.py.\n'
                ' * These are one-way hashes: the real codes cannot be read from this file.\n */\n'
                'window.STUDIOLO_LICENSE_DATA = {\n  batches: ' + body + '\n};\n')


def write_sheet(path, rows, batch_id):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = 'Licence codes'
    headers = ['#', 'Licence code', 'Batch', 'Given to', 'Email', 'Date given', 'Notes']
    ws.append(headers)
    for i, code in enumerate(rows, 1):
        ws.append([i, code, batch_id, '', '', '', ''])
    for col, width in zip('ABCDEFG', (7, 24, 10, 26, 30, 14, 34)):
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='EA580C')
        cell.alignment = Alignment(vertical='center')
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.font = Font(name='Consolas')
    ws.freeze_panes = 'A2'
    table = Table(displayName='Codes_' + batch_id, ref='A1:G%d' % (len(rows) + 1))
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    wb.save(path)
    return True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--count', type=int, default=1000)
    ap.add_argument('--out', required=True, help='folder OUTSIDE this project for the private code list')
    ap.add_argument('--batch-id', help='batch name (default B1, B2, …); DEV marks developer codes')
    args = ap.parse_args()
    out = os.path.abspath(args.out)
    if out == ROOT or out.startswith(ROOT + os.sep):
        raise SystemExit('Refusing to write codes inside the project. Choose a private folder elsewhere.')
    os.makedirs(out, exist_ok=True)

    batches = load_batches()
    batch_id = args.batch_id or 'B%d' % (sum(1 for b in batches if b['id'].startswith('B')) + 1)
    if any(b['id'] == batch_id for b in batches):
        raise SystemExit('Batch %s already exists.' % batch_id)
    salt = secrets.token_hex(16)
    codes = set()
    while len(codes) < args.count:
        codes.add(new_code())
    codes = sorted(codes, key=lambda c: secrets.randbits(32))
    batches.append({
        'id': batch_id,
        'created': datetime.date.today().isoformat(),
        'salt': salt,
        'hashes': sorted(fingerprint(salt, c) for c in codes)
    })
    save_batches(batches)

    label = 'Studiolo developer code' if batch_id == 'DEV' else 'Studiolo licence codes - batch %s' % batch_id
    stem = os.path.join(out, label)
    with open(stem + '.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['#', 'Licence code', 'Batch', 'Given to', 'Email', 'Date given', 'Notes'])
        for i, c in enumerate(codes, 1):
            w.writerow([i, c, batch_id, '', '', '', ''])
    has_xlsx = write_sheet(stem + '.xlsx', codes, batch_id)
    print('Batch %s: %d codes' % (batch_id, len(codes)))
    print('Private list:', stem + ('.xlsx and .csv' if has_xlsx else '.csv'))
    print('Fingerprints added to', os.path.relpath(CODES_JS, ROOT))


if __name__ == '__main__':
    main()
